import streamlit as st
import requests
import os
import sys
from pathlib import Path
import time
import io # Needed for handling file streams
import datetime # Import datetime for date handling
import json

# Import new libraries for file handling
import PyPDF2
import docx # python-docx
import pytesseract
from PIL import Image # Pillow

# Import OpenAI for AI interaction
from openai import OpenAI

# Import dotenv for loading environment variables
from dotenv import load_dotenv

# --- Configuration ---
PROJECT_ROOT = Path(__file__).resolve().parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.append(str(PROJECT_ROOT))

FASTAPI_URL = os.getenv("FASTAPI_URL", "http://127.0.0.1:8000")
EXTRACT_ENDPOINT = f"{FASTAPI_URL}/extract-context"

# Import database and generator functions
try:
    from database import SessionLocal, create_proposal, add_proposal_data, get_proposal_with_data, Proposal, init_db
    from proposal_generator import load_jinja_template, find_jinja_placeholders, fill_template_jinja, TEMPLATE_DIR, get_proposal_data_as_dict # <-- Re-add get_proposal_data_as_dict
except ImportError as e:
    st.error(f"Fatal Error: Could not import modules: {e}. Check paths and installations.")
    st.stop()
except Exception as e:
    st.error(f"Fatal Error: An unexpected error occurred during module import: {e}")
    st.stop()


# --- Text Extraction Helper Functions ---

def extract_text_from_txt(file_content: bytes) -> str:
    """Extracts text from bytes assuming UTF-8 encoding."""
    try:
        return file_content.decode('utf-8')
    except UnicodeDecodeError:
        st.warning("Could not decode .txt file as UTF-8, trying latin-1.")
        try:
            return file_content.decode('latin-1')
        except Exception as e:
            st.error(f"Error decoding .txt file: {e}")
            return ""

def extract_text_from_pdf(file_stream) -> str:
    """Extracts text from a PDF file stream."""
    try:
        reader = PyPDF2.PdfReader(file_stream)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        st.error(f"Error reading PDF content: {e}")
        return ""

def extract_text_from_docx(file_stream) -> str:
    """Extracts text from a DOCX file stream."""
    try:
        document = docx.Document(file_stream)
        text = "\n".join([para.text for para in document.paragraphs])
        return text
    except Exception as e:
        st.error(f"Error reading DOCX content: {e}")
        return ""

def extract_text_from_image(file_stream) -> str:
    """Extracts text from an image file stream using Tesseract OCR."""
    try:
        # Check if tesseract is installed and accessible
        # You might need to configure the path if tesseract isn't in the system PATH
        # pytesseract.pytesseract.tesseract_cmd = r'/path/to/tesseract' # Example if needed
        img = Image.open(file_stream)
        text = pytesseract.image_to_string(img)
        return text
    except Exception as e:
        # Catch specific Tesseract errors if possible, e.g., TesseractNotFoundError
        st.error(f"Error performing OCR on image: {e}. Ensure Tesseract is installed and in PATH.")
        return ""


# --- Other Helper Functions ---

def get_available_templates() -> list[str]:
    """Gets a list of template names (without .md extension)."""
    if not TEMPLATE_DIR.is_dir():
        st.warning(f"Template directory '{TEMPLATE_DIR}' not found.")
        return []
    templates = [f.stem for f in TEMPLATE_DIR.glob("*.md")]
    if not templates:
        st.warning(f"No '.md' templates found in '{TEMPLATE_DIR}'.")
    return templates


# --- Database Initialization ---
if 'db_initialized' not in st.session_state:
    st.session_state.db_initialized = False

def ensure_db_initialized():
    if not st.session_state.db_initialized:
        try:
            print("Attempting to initialize database...")
            init_db() # Initialize DB on first interaction if needed
            st.session_state.db_initialized = True
            print("Database initialized successfully.")
        except Exception as e:
            st.error(f"Failed to initialize database: {e}")
            st.stop()

ensure_db_initialized()


# --- Streamlit App ---
st.set_page_config(layout="wide")
st.title("AI Proposal Generator Assistant")

# --- State Initialization ---
# Keep existing state variables, add one for combined text
if 'proposal_id' not in st.session_state:
    st.session_state.proposal_id = None
if 'proposal_name' not in st.session_state:
    st.session_state.proposal_name = None
if 'extracted_data' not in st.session_state:
    st.session_state.extracted_data = None
# if 'uploaded_filename' not in st.session_state: # Replace single filename with list
#     st.session_state.uploaded_filename = None
if 'uploaded_file_names' not in st.session_state:
    st.session_state.uploaded_file_names = []
if 'combined_context_text' not in st.session_state:
    st.session_state.combined_context_text = ""
if 'template_name' not in st.session_state:
    st.session_state.template_name = None
if 'template_object' not in st.session_state: # Store Jinja template object
    st.session_state.template_object = None
if 'template_placeholders' not in st.session_state:
    st.session_state.template_placeholders = set()
if 'current_proposal_data' not in st.session_state:
    st.session_state.current_proposal_data = {}
if 'missing_keys' not in st.session_state:
    st.session_state.missing_keys = set()
if 'interview_data' not in st.session_state:
    st.session_state.interview_data = {}
if 'final_proposal_content' not in st.session_state:
    st.session_state.final_proposal_content = None


# --- Initialize OpenAI Client for Streamlit --- 
# This is needed for the AI pre-fill function
# Load env variables FOR the client init
load_dotenv()

openai_api_key_streamlit = os.getenv("OPENAI_API_KEY")
if not openai_api_key_streamlit:
    # Don't raise error here, just warn, as core extraction happens in backend
    st.warning("OPENAI_API_KEY not found for Streamlit app. AI pre-fill will be disabled.")
    client = None # Set client to None if key is missing
else:
    try:
        client = OpenAI(api_key=openai_api_key_streamlit)
    except Exception as client_err:
         st.error(f"Failed to initialize OpenAI client in Streamlit: {client_err}")
         client = None


# --- Workflow Steps ---

# Step 1: File Upload and Context Extraction
st.header("1. Upload Context & Extract Data")
st.write("Upload one or more context documents (.txt, .pdf, .docx, .png, .jpg, .jpeg). Their text content will be combined and sent to the AI for key information extraction.")

uploaded_files = st.file_uploader(
    "Upload context documents:",
    type=['txt', 'pdf', 'docx', 'png', 'jpg', 'jpeg'],
    accept_multiple_files=True, # Allow multiple files
    key="file_uploader"
)

# Check if the list of uploaded files has changed
current_file_names = sorted([f.name for f in uploaded_files])
if uploaded_files and current_file_names != st.session_state.get('uploaded_file_names', []):
    st.info(f"Processing {len(uploaded_files)} file(s): {', '.join(current_file_names)}")
    # Reset downstream state when new files are uploaded
    st.session_state.proposal_id = None
    st.session_state.proposal_name = None
    st.session_state.extracted_data = None
    st.session_state.current_proposal_data = {}
    st.session_state.interview_data = {}
    st.session_state.missing_keys = set()
    st.session_state.final_proposal_content = None
    st.session_state.combined_context_text = ""
    st.session_state.uploaded_file_names = current_file_names # Mark as processed

    all_text = []
    has_errors = False
    with st.spinner("Extracting text from uploaded files..."):
        for uploaded_file in uploaded_files:
            st.write(f" - Processing {uploaded_file.name}...")
            file_stream = io.BytesIO(uploaded_file.getvalue())
            text = ""
            file_type = Path(uploaded_file.name).suffix.lower()

            if file_type == '.txt':
                text = extract_text_from_txt(uploaded_file.getvalue()) # txt needs bytes directly
            elif file_type == '.pdf':
                text = extract_text_from_pdf(file_stream)
            elif file_type == '.docx':
                text = extract_text_from_docx(file_stream)
            elif file_type in ['.png', '.jpg', '.jpeg']:
                text = extract_text_from_image(file_stream)
            else:
                st.warning(f"Unsupported file type: {uploaded_file.name}")
                continue # Skip unsupported types

            if text:
                all_text.append(text)
            else:
                # Error messages are shown within the extractor functions
                st.warning(f"Could not extract text from {uploaded_file.name}.")
                # Decide if partial failure should stop the whole process
                # has_errors = True

    # Combine text from all successfully processed files
    st.session_state.combined_context_text = "\n\n--- File Separator ---\n\n".join(all_text)

    if not st.session_state.combined_context_text.strip():
        st.error("Failed to extract any text from the uploaded files.")
        st.session_state.uploaded_file_names = [] # Allow re-upload if all failed
    else:
        st.success("Text extraction complete.")
        st.text_area("Combined Text Preview (first 1000 chars):", st.session_state.combined_context_text[:1000] + "...", height=200, key="combined_preview")

        # --- Call AI for Extraction --- (Moved here to run after all files processed)
        with st.spinner("Sending combined text to AI for context extraction..."):
            try:
                response = requests.post(EXTRACT_ENDPOINT, json={"text": st.session_state.combined_context_text}, timeout=120) # Increased timeout
                response.raise_for_status()
                extracted = response.json().get("data", {})
                st.session_state.extracted_data = extracted if isinstance(extracted, dict) else {}
                st.success("AI extraction successful!")

                # --- Create Proposal in DB --- (Moved here)
                db_session = SessionLocal()
                try:
                    proposal_name_base = "proposal_from_upload"
                    if current_file_names:
                         proposal_name_base = Path(current_file_names[0]).stem # Use first filename as base
                    timestamp = int(time.time())
                    st.session_state.proposal_name = f"{proposal_name_base}_{timestamp}"
                    new_proposal = create_proposal(db_session, st.session_state.proposal_name)
                    st.session_state.proposal_id = new_proposal.id

                    if st.session_state.extracted_data:
                        add_proposal_data(db_session, new_proposal.id, st.session_state.extracted_data)
                        st.session_state.current_proposal_data = st.session_state.extracted_data
                        st.success(f"Extracted data saved to new proposal '{st.session_state.proposal_name}' (ID: {st.session_state.proposal_id}).")
                    else:
                        st.warning("AI did not return any structured data.")
                        st.session_state.current_proposal_data = {}
                    st.rerun() # Rerun to update UI state

                except Exception as db_error:
                    st.error(f"Database error after extraction: {db_error}")
                    st.session_state.proposal_id = None
                    st.session_state.proposal_name = None
                    st.session_state.current_proposal_data = {}
                finally:
                    db_session.close()

            except requests.exceptions.Timeout:
                 st.error("Error calling extraction API: Request timed out.")
            except requests.exceptions.ConnectionError:
                 st.error(f"Error calling extraction API: Could not connect to the backend at {FASTAPI_URL}.")
            except requests.exceptions.RequestException as e:
                 st.error(f"Error calling extraction API: {e}")
            except Exception as e:
                 st.error(f"An unexpected error occurred during AI processing: {e}")

            # If extraction or DB saving failed, clear related state
            if st.session_state.proposal_id is None:
                st.session_state.extracted_data = None
                st.session_state.current_proposal_data = {}

# Step 2: Review Data and Select Template
if st.session_state.proposal_id is not None:
    st.header("2. Review Data & Select Template")
    st.info(f"Working with Proposal: **{st.session_state.proposal_name}** (ID: {st.session_state.proposal_id})")

    st.subheader("Current Proposal Data")
    if st.session_state.current_proposal_data:
         st.json(st.session_state.current_proposal_data)
    else:
         st.write("No data extracted or added yet.")

    st.subheader("Select Template")
    available_templates = get_available_templates()
    if not available_templates:
        st.error("Cannot proceed without templates. Please add '.md' files to the 'templates' directory.")
    else:
        current_template_index = None
        if st.session_state.template_name in available_templates:
             current_template_index = available_templates.index(st.session_state.template_name)

        selected_template = st.selectbox(
            "Select a proposal template:",
            options=available_templates,
            index=current_template_index if current_template_index is not None else 0,
            key="template_selector"
        )

        # Load Jinja template object if selection changes or not yet loaded
        if selected_template and (selected_template != st.session_state.template_name or not st.session_state.template_object):
            st.session_state.template_name = selected_template
            # Load Jinja Template Object
            st.session_state.template_object = load_jinja_template(selected_template)
            if st.session_state.template_object:
                # Get placeholders using Jinja parser
                st.session_state.template_placeholders = find_jinja_placeholders(f"{selected_template}.md")
                # Get template content for preview (can read source from loaded template)
                try:
                    template_content_preview = st.session_state.template_object.environment.loader.get_source(st.session_state.template_object.environment, st.session_state.template_object.name)[0]
                    st.text_area("Template Preview:", template_content_preview, height=150, key="template_preview")
                except Exception:
                    st.warning("Could not display template preview.") # Non-critical
                
                st.session_state.missing_keys = set()
                st.session_state.interview_data = {}
                st.session_state.final_proposal_content = None
                st.rerun()
            else:
                st.error(f"Failed to load template '{selected_template}'.")
                st.session_state.template_object = None
                st.session_state.template_placeholders = set()


# --- AI Pre-fill Helper Function ---

def get_ai_best_guesses(client: OpenAI, context: str, existing_data: dict, missing_keys: list[str]) -> dict:
    """Makes a second AI call to get best guesses for missing keys."""
    if not missing_keys:
        return {}
    
    keys_to_guess_str = ", ".join(missing_keys)
    existing_data_str = json.dumps(existing_data, indent=2)
    
    system_prompt = (
        "You are an expert assistant helping pre-fill a proposal template. "
        "Based on the provided context, existing extracted data, and common business practices, "
        "provide reasonable default values or best guesses for the listed missing keys. "
        "Output MUST be a valid JSON object containing ONLY the keys requested with their guessed values. "
        "Do not include keys that were already extracted. Use standard formats (e.g., dates as YYYY-MM-DD, percentages as numbers)."
    )
    user_prompt = (
        f"Context: ```{context}```\n\n"
        f"Existing Data: ```{existing_data_str}```\n\n"
        f"Please provide best guesses ONLY for these missing keys: {keys_to_guess_str}"
    )

    st.info(f"Asking AI for best guesses for: {keys_to_guess_str}")
    with st.spinner("AI is generating best guesses for missing fields..."):
        try:
            # Use the passed client object
            response = client.chat.completions.create(
                model=os.getenv("OPENAI_MODEL", "gpt-4o"), 
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt}
                ],
                temperature=0.3, # Allow a bit more creativity for guesses
                response_format={ "type": "json_object" } 
            )
            guessed_data = json.loads(response.choices[0].message.content)
            st.success("AI provided best guesses.")
            # Basic validation: ensure it only returned requested keys
            valid_guesses = {k: v for k, v in guessed_data.items() if k in missing_keys}
            if len(valid_guesses) < len(guessed_data):
                st.warning("AI returned some unexpected keys in its guesses.")
            return valid_guesses
        except Exception as e:
            st.error(f"Error getting AI best guesses: {e}")
            return {}


# Helper for date calculations
def get_next_weekday(start_date, weekday):
    """Find the next specific weekday (0=Mon, 6=Sun) on or after start_date."""
    days_ahead = weekday - start_date.weekday()
    if days_ahead < 0: # Target day already happened this week
        days_ahead += 7
    return start_date + datetime.timedelta(days=days_ahead)

def add_business_days(start_date, days_to_add):
    """Add business days, skipping weekends."""
    # This is a simple implementation; edge cases like holidays aren't handled.
    current_date = start_date
    while days_to_add > 0:
        current_date += datetime.timedelta(days=1)
        weekday = current_date.weekday()
        if weekday < 5: # Monday to Friday
            days_to_add -= 1
    # Ensure final date is not on weekend (adjust forward)
    while current_date.weekday() >= 5:
        current_date += datetime.timedelta(days=1)
    return current_date

def calculate_default_dates():
    """Calculates default dates based on business logic."""
    today = datetime.date.today()
    defaults = {}
    try:
        # Proposal Date: Today
        defaults['proposal_date'] = today.strftime("%B %d, %Y")

        # Contract Date: Friday of the *following* week
        contract_date_obj = get_next_weekday(today + datetime.timedelta(days=1), 4) # Find next Fri
        if contract_date_obj.weekday() == 4 and contract_date_obj <= today + datetime.timedelta(days=7): # If it's this week's Friday or past
            contract_date_obj = contract_date_obj + datetime.timedelta(days=7) # Move to next week's Friday
        defaults['contract_date'] = contract_date_obj.strftime("%B %d, %Y")

        # Advertising Start Date: Second Monday after Contract Date
        adv_start_monday_1 = get_next_weekday(contract_date_obj, 0) # First Monday on or after contract date
        adv_start_monday_2 = adv_start_monday_1 + datetime.timedelta(days=7) # Second Monday
        defaults['advertising_start_date'] = adv_start_monday_2.strftime("%B %d, %Y")

        # Auction End Date: Wednesday 3 full weeks after Advertising Start Date
        adv_start_date_obj = adv_start_monday_2
        auction_end_wed = get_next_weekday(adv_start_date_obj + datetime.timedelta(days=21), 2) # Find Wed >= 3 weeks later
        defaults['auction_end_date'] = auction_end_wed.strftime("%B %d, %Y")

        # Closing Date: 30 calendar days after Auction End Date, adjusted to next business day
        auction_end_date_obj = auction_end_wed
        closing_date_initial = auction_end_date_obj + datetime.timedelta(days=30)
        # Adjust if weekend
        while closing_date_initial.weekday() >= 5: # 5=Sat, 6=Sun
            closing_date_initial += datetime.timedelta(days=1)
        defaults['closing_date'] = closing_date_initial.strftime("%B %d, %Y")
    except Exception as date_err:
        print(f"Error calculating default dates: {date_err}")
        # Return potentially partial defaults
    return defaults

DEFAULT_MARKETING_COSTS = {
    'marketing_drone_cost': 500.00,
    'marketing_facebook_cost': 1000.00,
    'marketing_google_cost': 1000.00,
    'marketing_mail_cost': 500.00,
    'marketing_sign_cost': 250.00
}

# Step 3: Interview / Review & Edit Mode
if st.session_state.template_object and st.session_state.proposal_id is not None:

    # --- Attempt AI Pre-fill --- 
    if 'ran_ai_prefill' not in st.session_state:
         st.session_state.ran_ai_prefill = False
         # Reset flag if template changes or new proposal loaded
         # This might need adjustment depending on exact desired flow

    # Calculate initially missing keys
    initial_missing_keys = st.session_state.template_placeholders - set(st.session_state.current_proposal_data.keys())
    calculated_keys = {'marketing_total_budget', 'retainer_amount'}
    initial_missing_keys.difference_update(calculated_keys)

    # Run AI pre-fill if missing keys exist and haven't run it yet for this cycle
    if initial_missing_keys and not st.session_state.ran_ai_prefill:
        # ... (AI guess logic remains the same - calls get_ai_best_guesses)
        if client:
             try:
                 best_guesses = get_ai_best_guesses(
                    client, 
                    st.session_state.combined_context_text,
                    st.session_state.current_proposal_data,
                    list(initial_missing_keys)
                 )
             except NameError: 
                  st.error("OpenAI client not initialized. Cannot get AI guesses.")
                  best_guesses = {}
        else:
             st.warning("AI Pre-fill skipped: OpenAI client not available.")
             best_guesses = {} 
        
        if best_guesses:
            # Merge AI guesses into current data
            st.session_state.current_proposal_data.update(best_guesses)
            # Save merged data (guesses only) to DB
            db_session = SessionLocal()
            try:
                save_data = {k: str(v) for k, v in best_guesses.items()}
                add_proposal_data(db_session, st.session_state.proposal_id, save_data)
                st.success("AI best guesses saved to database.")
            except Exception as e:
                st.error(f"Failed to save AI best guesses to database: {e}")
            finally:
                db_session.close()
                
        st.session_state.ran_ai_prefill = True # Mark pre-fill as done
        # Don't rerun here, proceed to apply defaults and show form

    # --- Apply Calculated Defaults --- 
    # Calculate defaults regardless of pre-fill run
    default_dates = calculate_default_dates()
    combined_defaults = {**default_dates, **DEFAULT_MARKETING_COSTS}
    
    # Apply defaults only if the key is STILL missing after initial load + AI guess
    applied_defaults_count = 0
    for key, default_value in combined_defaults.items():
        if key not in st.session_state.current_proposal_data or not st.session_state.current_proposal_data[key]:
            st.session_state.current_proposal_data[key] = default_value
            applied_defaults_count += 1
            print(f"Applied default for missing key '{key}': {default_value}")
    
    if applied_defaults_count > 0:
        # Optionally save these defaults back to DB immediately
        # Or wait until user submits the review form
        pass # For now, just update session state

    # --- Display Review/Edit Form --- 
    st.header("3. Review and Edit Data")
    st.write("Review the data extracted and guessed by the AI, along with calculated defaults. Make any necessary corrections before generating the proposal.")

    # Determine keys required by template, excluding calculated ones
    display_keys = st.session_state.template_placeholders - calculated_keys

    with st.form("review_edit_form"):
        edited_responses = {}
        date_keys = {'proposal_date', 'auction_end_date', 'closing_date', 'contract_date', 'advertising_start_date'}
        cost_keys = DEFAULT_MARKETING_COSTS.keys() # Get keys for cost inputs

        for key in sorted(list(display_keys)):
            # Get the current value from session state (includes extraction, guesses, defaults)
            current_value = st.session_state.current_proposal_data.get(key, "")
            
            if key in date_keys:
                # Parse current string value to date obj for widget
                default_date = None
                if isinstance(current_value, str) and current_value:
                    try: # Try preferred format first
                        default_date = datetime.datetime.strptime(current_value, "%B %d, %Y").date()
                    except ValueError: # Try AI guess format
                        try:
                            default_date = datetime.datetime.strptime(current_value, "%Y-%m-%d").date()
                        except ValueError:
                             print(f"Warning: Could not parse date string '{current_value}' for key '{key}'. Defaulting widget to today.")
                             default_date = datetime.date.today()
                elif isinstance(current_value, datetime.date):
                     default_date = current_value # Use if already date obj
                else:
                   default_date = datetime.date.today() # Default if empty or other type
                
                edited_responses[key] = st.date_input(
                    f"{key.replace('_', ' ').title()}", 
                    value=default_date, 
                    key=f"review_{key}"
                )
            elif key in cost_keys:
                 # Use number input for costs
                 default_cost = 0.0
                 try:
                     default_cost = float(str(current_value).replace('$','').replace(',','') or 0.0)
                 except ValueError:
                      print(f"Warning: Could not parse cost value '{current_value}' for key '{key}'. Defaulting widget to 0.")
                 edited_responses[key] = st.number_input(
                     f"{key.replace('_', ' ').title()}",
                     value=default_cost,
                     format="%.2f",
                     key=f"review_{key}"
                 )
            else: # Standard text input for others
                edited_responses[key] = st.text_input(f"{key.replace('_', ' ').title()}", value=str(current_value), key=f"review_{key}")

        submitted = st.form_submit_button("Save Changes and Proceed")
        if submitted:
            valid_submission = True
            final_edited_data = {}
            for key, value in edited_responses.items():
                # Format dates before saving, convert others to string
                if key in date_keys and isinstance(value, (datetime.date, datetime.datetime)):
                    cleaned_value = value.strftime("%B %d, %Y") 
                else:
                    # Convert number inputs back to string for DB consistency, or keep as number if DB supports it
                    cleaned_value = str(value).strip()
                    
                if not cleaned_value:
                    st.error(f"Value for '{key}' cannot be empty.")
                    valid_submission = False
                final_edited_data[key] = cleaned_value

            if valid_submission:
                # Save the potentially edited data back to DB
                db_session = SessionLocal()
                try:
                    add_proposal_data(db_session, st.session_state.proposal_id, final_edited_data)
                    st.success("Changes saved to database.")
                    # Update session state with the final edited values
                    st.session_state.current_proposal_data.update(final_edited_data)
                    st.session_state.missing_keys = set() # Clear missing keys as form was submitted
                    # Reset prefill flag for next potential run?
                    st.session_state.ran_ai_prefill = False 
                    st.rerun() # Rerun to proceed to Step 4 or refresh form if needed
                except Exception as e:
                    st.error(f"Failed to save changes to database: {e}")
                finally:
                    db_session.close()
            # No else needed - form just redisplays with entered values if invalid

# Step 4: Generate and Download Proposal
# Only show if proposal exists and form was successfully submitted (implies no missing keys relevant to template)
# We can check if final_edited_data exists from a successful submit, or simplify check
can_generate = (
    st.session_state.template_object and
    st.session_state.proposal_id is not None and
    # Check if the review form was submitted successfully OR if initially no keys were missing
    # A simple check could be if missing_keys is empty AFTER the review form logic
    not st.session_state.missing_keys # Assuming successful submit clears this
)


if can_generate:
    st.header("4. Generate & Download Proposal")
    st.success("All necessary information is available.")

    if st.button("Generate Proposal Document", key="generate_button"):
        db_session = SessionLocal()
        try:
            # Fetch latest data and ensure number conversion for Jinja context
            db_proposal = get_proposal_with_data(db_session, st.session_state.proposal_id)
            full_context = get_proposal_data_as_dict(db_proposal) # Use helper that converts numbers
            
            # Add calculated values
            marketing_keys = ['marketing_facebook_cost', 'marketing_google_cost', 'marketing_mail_cost', 'marketing_drone_cost', 'marketing_sign_cost']
            marketing_total_budget = sum(float(full_context.get(key, 0) or 0) for key in marketing_keys)
            retainer_amount = 10000 - marketing_total_budget
            full_context['marketing_total_budget'] = marketing_total_budget
            full_context['retainer_amount'] = retainer_amount
            
            # --- Format dates in context before rendering ---
            date_keys_to_format = {'proposal_date', 'auction_end_date', 'closing_date', 'contract_date', 'advertising_start_date'}
            for key in date_keys_to_format:
                if key in full_context:
                    try:
                        # Attempt to parse the date (assuming YYYY-MM-DD or other common formats)
                        date_obj = None
                        try: # Try YYYY-MM-DD first (common AI output)
                           date_obj = datetime.datetime.strptime(str(full_context[key]), "%Y-%m-%d").date()
                        except ValueError: # Try Month D, YYYY (format we save from date picker)
                           try:
                               date_obj = datetime.datetime.strptime(str(full_context[key]), "%B %d, %Y").date()
                           except ValueError:
                               pass # Add more formats if needed
                        
                        # If parsing succeeded, format it
                        if date_obj:
                           full_context[key] = date_obj.strftime("%B %d, %Y") 
                        # else: keep original string if parsing failed
                    except Exception as fmt_e: # Catch any unexpected errors during formatting
                         print(f"Warning: Could not parse/format date for key '{key}': {fmt_e}")
                         # Keep original string value
            
            # Render using Jinja template object
            st.session_state.final_proposal_content = fill_template_jinja(st.session_state.template_object, full_context)

            if st.session_state.final_proposal_content.startswith("Error rendering template:"):
                st.error("Proposal generation failed due to template rendering error.")
                st.session_state.final_proposal_content = None # Clear on error
            else:
                st.success("Proposal content generated successfully!")
                st.markdown("### Generated Proposal Preview:")
                st.text_area("Generated Proposal Preview Text", st.session_state.final_proposal_content, height=300, key="final_preview", label_visibility="collapsed")
        except Exception as e:
            st.error(f"An error occurred during final proposal generation: {e}")
            st.session_state.final_proposal_content = None
        finally:
            db_session.close()

    # ... (Download Button logic - remains the same) ...
    if st.session_state.final_proposal_content:
        proposal_filename = f"{st.session_state.proposal_name or 'proposal'}.md"
        st.download_button(
            label="Download Proposal as Markdown",
            data=st.session_state.final_proposal_content,
            file_name=proposal_filename,
            mime="text/markdown",
            key="download_button"
        )

# --- Footer/Debug ---
# with st.expander("Debug Info (Session State)"):
#     st.write(st.session_state) 